"""Update the user's Excel logbook from the logbook database (B).
Only the flight lines are filled (month, day, type, registration, crew, route, remarks and the time
columns); formulas, totals, merged cells and formatting are left alone, and a line that already holds
data is never overwritten. Adapted from the update step of the earlier Excel-based application."""
import re
import secrets
import sqlite3
import time
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

from flask import Blueprint, after_this_request, flash, g, redirect, render_template, request, send_file, url_for
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from .auth import login_required
from .errors import safe_message
from .seclog import log_event
from .uploads import check_xlsx
from .db import logbook_connect, user_file

export_bp = Blueprint("export", __name__, url_prefix="/export")

MAX_UPLOAD = 25 * 1024 * 1024
LINES_PER_PAGE = 30
COL_MONTH, COL_DAY = 2, 3
TEXT_COLUMNS = {"aircraft_type": 4, "registration": 5, "pilot_in_command": 6, "copilot": 7,
                "departure": 8, "arrival": 9, "remarks": 10}
TIME_COLUMNS = {
    "single_engine_dual_day": 12, "single_engine_pic_day": 13, "single_engine_dual_night": 14,
    "single_engine_pic_night": 15, "multi_engine_dual_day": 18, "multi_engine_pic_day": 19,
    "multi_engine_copi_day": 20, "multi_engine_dual_night": 21, "multi_engine_pic_night": 22,
    "multi_engine_copi_night": 23, "ifr": 24, "hood": 25, "ftd": 26, "ifr_approach": 27,
    "cross_country_dual_day": 28, "cross_country_pic_day": 29, "cross_country_dual_night": 30,
    "cross_country_pic_night": 31, "landings_day": 32, "landings_night": 33,
    "instruction_day": 34, "instruction_night": 35}
INPUT_COLUMNS = list(range(2, 11)) + list(range(12, 36))     # column K is half of the merged remarks cell
FLIGHT_NUM_RE = re.compile(r"([A-Z]{2,3}\d{2,5})")


def match_token(remarks):
    """Flight number (ex. AC8942) found in the remarks, else the remarks without punctuation."""
    remarks = (remarks or "").strip().upper()
    m = FLIGHT_NUM_RE.search(remarks)
    return m.group(1) if m else re.sub(r"[^A-Z0-9]+", " ", remarks).strip()


def two_digits(value):
    try:
        return f"{int(float(str(value))):02d}"
    except (TypeError, ValueError):
        return str(value or "").strip()


def _reg(value):
    return re.sub(r"[^A-Z0-9]", "", str(value or "").upper())


def flight_rows(ws):
    """Flight lines of a year sheet: the 30 lines above each 'GRAND TOTAL' row, page after page."""
    rows = []
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 8).value or "").strip().upper() == "GRAND TOTAL" and r - LINES_PER_PAGE >= 1:
            rows.extend(range(r - LINES_PER_PAGE, r))
    return rows


def row_is_writable(ws, r):
    """False when a cell we would write to is the hidden part of a merged range (read-only in openpyxl)."""
    return not any(isinstance(ws.cell(r, c), MergedCell) for c in INPUT_COLUMNS)


def line_is_empty(ws, r):
    return all(ws.cell(r, c).value in (None, "") for c in INPUT_COLUMNS)


def number(v):
    try:
        v = float(v)
    except (TypeError, ValueError):
        return None
    if v == 0:
        return None
    return int(v) if v == int(v) else v


def load_flights(rows):
    """Logbook rows with a valid date, grouped by year, in the logbook order."""
    by_year, invalid = defaultdict(list), 0
    for r in rows:
        try:
            d = date.fromisoformat(str(r.get("date") or "")[:10])
        except ValueError:
            invalid += 1
            continue
        by_year[d.year].append({**r, "_date": d, "_key": (f"{d.month:02d}", f"{d.day:02d}", match_token(r.get("remarks")))})
    return by_year, invalid


def update_workbook(wb, rows):
    """Fill the free flight lines of every year sheet with the flights the sheet does not hold yet."""
    by_year, invalid = load_flights(rows)
    wb.calculation.fullCalcOnLoad = True            # Excel recomputes every formula when the file is opened
    report = {"sheets": [], "no_sheet": {}, "invalid_dates": invalid}
    for ws in wb.worksheets:
        if not re.fullmatch(r"\d{4}", ws.title):
            continue
        year = int(ws.title)
        all_slots = flight_rows(ws)
        slots = [r for r in all_slots if row_is_writable(ws, r)]      # lines with merged cells are left alone
        existing, last_date = defaultdict(list), None
        for r in slots:
            if not line_is_empty(ws, r):
                key = (two_digits(ws.cell(r, COL_MONTH).value), two_digits(ws.cell(r, COL_DAY).value),
                       match_token(ws.cell(r, TEXT_COLUMNS["remarks"]).value))
                existing[key].append(_reg(ws.cell(r, TEXT_COLUMNS["registration"]).value))
                try:
                    last_date = date(year, int(key[0]), int(key[1]))
                except ValueError:
                    pass
        filled = sum(len(v) for v in existing.values())
        pending, present = [], 0
        for f in by_year.get(year, []):
            lines, reg = existing.get(f["_key"], []), _reg(f.get("registration"))
            idx = next((i for i, x in enumerate(lines) if reg and x == reg), None)          # même appareil d'abord
            if idx is None:
                idx = next((i for i, x in enumerate(lines) if not reg or not x), None)   # sinon ligne sans immatriculation
            if idx is not None:
                lines.pop(idx)
                present += 1
            else:
                pending.append(f)
        free = [r for r in slots if line_is_empty(ws, r)]
        added = pending[:len(free)]
        for r, f in zip(free, added):
            ws.cell(r, COL_MONTH).value, ws.cell(r, COL_DAY).value = f["_key"][0], f["_key"][1]
            for column, c in TEXT_COLUMNS.items():
                text = str(f.get(column) or "").strip()
                if text:
                    ws.cell(r, c).value = text
                    if text.startswith(("=", "+", "-", "@")):
                        ws.cell(r, c).data_type = "s"      # jamais interprété comme une formule
            for column, c in TIME_COLUMNS.items():
                value = number(f.get(column))
                if value is not None:
                    ws.cell(r, c).value = value
        report["sheets"].append({
            "sheet": ws.title, "slots": len(slots), "merged_skipped": len(all_slots) - len(slots), "filled": filled, "in_db": len(by_year.get(year, [])),
            "present": present, "added": len(added), "no_slot": len(pending) - len(added),
            "out_of_order": sum(1 for f in added if last_date and f["_date"] < last_date),
            "not_in_db": sum(len(v) for v in existing.values()),
            "samples": [f"{f['_key'][1]}/{f['_key'][0]} {str(f.get('remarks') or '')[:40]}" for f in added[:8]]})
    sheet_years = {int(ws.title) for ws in wb.worksheets if re.fullmatch(r"\d{4}", ws.title)}
    report["no_sheet"] = {y: len(v) for y, v in sorted(by_year.items()) if y not in sheet_years}
    report["added_total"] = sum(s["added"] for s in report["sheets"])
    return report


def _tmp_dir():
    path = Path(user_file("export_tmp", g.username))
    path.mkdir(parents=True, exist_ok=True)
    for old in path.glob("*.xlsx"):                 # temporary copies never stay more than an hour
        if time.time() - old.stat().st_mtime > 3600:
            old.unlink(missing_ok=True)
    return path


def _load_rows():
    try:
        conn = logbook_connect()
        conn.row_factory = sqlite3.Row
        rows = [dict(r) for r in conn.execute("SELECT * FROM logbook ORDER BY id")]
        conn.close()
    except sqlite3.Error:
        rows = []
    return rows


@export_bp.route("/", methods=["GET", "POST"])
@login_required
def index():
    if request.method == "GET":
        return render_template("export.html", report=None, token=None, filename=None)
    upload = request.files.get("logbook")
    if not upload or not upload.filename.lower().endswith(".xlsx"):
        flash("Choisis un fichier Excel au format .xlsx.")
        return redirect(url_for("export.index"))
    folder = _tmp_dir()
    source = folder / f"{secrets.token_hex(8)}_source.xlsx"
    upload.save(source)
    try:
        if source.stat().st_size > MAX_UPLOAD:
            raise ValueError("Fichier trop volumineux (25 Mo maximum).")
        check_xlsx(source)
        wb = load_workbook(source)
        report = update_workbook(wb, _load_rows())
        token = secrets.token_urlsafe(16)
        wb.save(folder / f"{token}.xlsx")
    except Exception as exc:                        # unreadable or unexpected workbook
        flash(f"Impossible de traiter ce fichier : {safe_message(exc)}")
        return redirect(url_for("export.index"))
    finally:
        source.unlink(missing_ok=True)              # the uploaded copy is never kept
    stem = re.sub(r"[^\w\- ]+", "", Path(upload.filename).stem).strip() or "logbook"
    return render_template("export.html", report=report, token=token, filename=f"{stem} - mis à jour.xlsx")


@export_bp.route("/download/<token>")
@login_required
def download(token):
    path = _tmp_dir() / f"{re.sub(r'[^A-Za-z0-9_-]', '', token)}.xlsx"
    if not path.exists():
        flash("Ce fichier n'est plus disponible : relance l'export.")
        return redirect(url_for("export.index"))

    @after_this_request
    def cleanup(response):
        path.unlink(missing_ok=True)                # the copy is removed once downloaded
        return response
    name = re.sub(r"[^\w\-. ]+", "", request.args.get("name") or "")[:100] or "logbook - mis à jour.xlsx"
    if not name.lower().endswith(".xlsx"):
        name += ".xlsx"
    log_event("data_export", target="logbook-excel-mis-a-jour")
    return send_file(path, as_attachment=True, download_name=name,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
