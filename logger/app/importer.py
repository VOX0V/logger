import io
import re
from openpyxl import load_workbook


def normalize_header(value):
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = re.sub(r"\s+", " ", text)
    return text


def read_excel_files(files):
    """Return rows from every sheet of every uploaded Excel workbook."""
    result = []
    for filename, data in files:
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise ValueError(f'Le fichier "{filename}" est illisible ou corrompu.') from exc
        for sheet in wb.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [normalize_header(v) for v in rows[0]]
            if not any(headers):
                continue
            result.append((filename, sheet.title, headers, rows[1:]))
    if not result:
        raise ValueError("Aucune feuille Excel exploitable n'a été trouvée.")
    return result


def build_records(sources, categories, rules_by_category):
    records = []
    report = []
    for filename, sheet, headers, rows in sources:
        header_indexes = {}
        for i, header in enumerate(headers):
            if header and header not in header_indexes:
                header_indexes[header] = i

        mappings = {}
        for cat in categories:
            rules = rules_by_category.get(cat["id"], [])
            match = next((rule for rule in rules if normalize_header(rule) in header_indexes), None)
            if match is not None:
                mappings[cat["column_name"]] = header_indexes[normalize_header(match)]

        matched_count = 0
        for row in rows:
            record = {}
            has_value = False
            for column, index in mappings.items():
                value = row[index] if index < len(row) else None
                if value is not None and str(value).strip() != "":
                    has_value = True
                record[column] = value
            if has_value:
                record["_import_source"] = filename
                records.append(record)
                matched_count += 1
        report.append({"filename": filename, "sheet": sheet, "matched": list(mappings.keys()), "rows": matched_count})
    return records, report
