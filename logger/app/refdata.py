"""Shared reference databases: airports (airport.db) and aircraft (aircrafts.db)."""
from pathlib import Path
from flask import current_app
import io
import re, sqlite3, unicodedata
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from .db import airport_connect, aircrafts_connect
from .importer import _read_table, FileImportError

ENGINES = ("single", "multi")
ENGINE_LABELS = {"single": "Monomoteur", "multi": "Multimoteur"}
CATEGORIES = (("airplane", "Avion"), ("helicopter", "Hélicoptère"), ("simulator", "Simulateur"))
ENGINE_TYPES = (("piston", "Piston"), ("turboprop", "Turbopropulseur"), ("turbine", "Turbine (hélicoptère)"),
                ("turbofan", "Turbofan"))

# Starting list of aircraft types (inserted once, when the table is created). To be checked by the user.
SEED_TYPES = [
    ("RH22", "Robinson", "R22", "helicopter", "single", "piston"),
    ("RH44", "Robinson", "R44", "helicopter", "single", "piston"),
    ("RH66", "Robinson", "R66", "helicopter", "single", "turbine"),
    ("AS350", "Airbus Helicopters", "AS350 Écureuil", "helicopter", "single", "turbine"),
    ("EC120", "Airbus Helicopters", "EC120 Colibri", "helicopter", "single", "turbine"),
    ("B505", "Bell", "505 Jet Ranger X", "helicopter", "single", "turbine"),
    ("BH06", "Bell", "206 JetRanger", "helicopter", "single", "turbine"),
    ("C150", "Cessna", "150", "airplane", "single", "piston"),
    ("C152", "Cessna", "152", "airplane", "single", "piston"),
    ("A152", "Cessna", "A152 Aerobat", "airplane", "single", "piston"),
    ("C172", "Cessna", "172", "airplane", "single", "piston"),
    ("TB10", "Socata", "TB 10 Tobago", "airplane", "single", "piston"),
    ("PA28", "Piper", "PA-28", "airplane", "single", "piston"),
    ("PA30", "Piper", "PA-30 Twin Comanche", "airplane", "multi", "piston"),
    ("PA31", "Piper", "PA-31 Navajo", "airplane", "multi", "piston"),
    ("CRJ", "Bombardier", "CRJ", "airplane", "multi", "turbofan"),
    ("SIM", None, "Simulateur", "simulator", None, None),
]


def _clean(value):
    text = "" if value is None else str(value).strip()
    return text or None


def _float(value, label, lo, hi, required=True):
    if value in (None, ""):
        if required:
            raise ValueError(f"{label} obligatoire.")
        return None
    try:
        number = float(str(value).replace(",", "."))
    except ValueError:
        raise ValueError(f"{label} invalide.")
    if not lo <= number <= hi:
        raise ValueError(f"{label} hors limites ({lo} à {hi}).")
    return number


# ---------------- airports ----------------

def init_airport_db():
    conn = airport_connect()
    conn.execute("""CREATE TABLE IF NOT EXISTS airports (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        icao_code TEXT UNIQUE,
        iata_code TEXT UNIQUE,
        full_name TEXT,
        country TEXT,
        latitude REAL NOT NULL,
        longitude REAL NOT NULL,
        elevation REAL,
        timezone TEXT,
        notes TEXT
    )""")
    empty = conn.execute("SELECT COUNT(*) FROM airports").fetchone()[0] == 0
    conn.commit(); conn.close()
    if empty:
        _import_default_airports()


def _import_default_airports():
    source = Path(current_app.root_path).parent / "defaults" / "airports.xlsx"
    if not source.exists():
        return
    wb = load_workbook(source, read_only=True, data_only=True)
    ws = wb["airports"] if "airports" in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return
    header = [str(h or "").strip() for h in rows[0]]
    conn = airport_connect()
    for row in rows[1:]:
        d = dict(zip(header, row))
        try:
            lat = _float(d.get("latitude"), "Latitude", -90, 90)
            lon = _float(d.get("longitude"), "Longitude", -180, 180)
        except ValueError:
            continue
        icao, iata = _clean(d.get("icao_code")), _clean(d.get("iata_code"))
        if not (icao or iata):
            continue
        conn.execute(
            "INSERT OR IGNORE INTO airports (icao_code, iata_code, full_name, country, latitude, longitude, elevation, timezone, notes) VALUES (?,?,?,?,?,?,?,?,?)",
            (icao.upper() if icao else None, iata.upper() if iata else None, _clean(d.get("full_name")), _clean(d.get("country")),
             lat, lon, _float(d.get("elevation"), "Altitude", -1000, 20000, required=False), _clean(d.get("timezone")), _clean(d.get("notes"))))
    conn.commit(); conn.close()


def find_airport(code):
    """Look an airport up by IATA code first, then ICAO. Returns a Row or None."""
    code = (code or "").strip().upper()
    if not code:
        return None
    conn = airport_connect()
    row = conn.execute("SELECT * FROM airports WHERE iata_code=?", (code,)).fetchone() \
        or conn.execute("SELECT * FROM airports WHERE icao_code=?", (code,)).fetchone()
    conn.close()
    return row


def list_airports():
    conn = airport_connect()
    rows = conn.execute("SELECT * FROM airports ORDER BY COALESCE(icao_code, iata_code)").fetchall()
    conn.close(); return rows


def get_airport(airport_id):
    conn = airport_connect()
    row = conn.execute("SELECT * FROM airports WHERE id=?", (airport_id,)).fetchone()
    conn.close(); return row


def save_airport(airport_id, form):
    icao = (_clean(form.get("icao_code")) or "").upper() or None
    iata = (_clean(form.get("iata_code")) or "").upper() or None
    if not (icao or iata):
        raise ValueError("Un code ICAO ou IATA est obligatoire.")
    lat = _float(form.get("latitude"), "Latitude", -90, 90)
    lon = _float(form.get("longitude"), "Longitude", -180, 180)
    elev = _float(form.get("elevation"), "Altitude", -1000, 20000, required=False)
    values = (icao, iata, _clean(form.get("full_name")), _clean(form.get("country")), lat, lon, elev,
              _clean(form.get("timezone")), _clean(form.get("notes")))
    conn = airport_connect()
    try:
        if airport_id:
            conn.execute("UPDATE airports SET icao_code=?, iata_code=?, full_name=?, country=?, latitude=?, longitude=?, elevation=?, timezone=?, notes=? WHERE id=?", values + (airport_id,))
        else:
            conn.execute("INSERT INTO airports (icao_code, iata_code, full_name, country, latitude, longitude, elevation, timezone, notes) VALUES (?,?,?,?,?,?,?,?,?)", values)
        conn.commit()
    except Exception as exc:
        conn.rollback()
        if "UNIQUE" in str(exc):
            raise ValueError("Ce code ICAO ou IATA existe déjà.")
        raise
    finally:
        conn.close()


def delete_airport(airport_id):
    conn = airport_connect()
    conn.execute("DELETE FROM airports WHERE id=?", (airport_id,))
    conn.commit(); conn.close()


# ---------------- aircraft ----------------

def init_aircrafts_db():
    conn = aircrafts_connect()
    conn.execute("""CREATE TABLE IF NOT EXISTS aircrafts (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        registration TEXT NOT NULL UNIQUE,
        type TEXT,
        engine TEXT CHECK (engine IN ('single','multi')),
        notes TEXT
    )""")   # legacy table by registration: kept as is, no longer used
    existed = conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='aircraft_types'").fetchone()
    conn.execute("""CREATE TABLE IF NOT EXISTS aircraft_types (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        code TEXT NOT NULL UNIQUE,
        make TEXT,
        model TEXT,
        category TEXT CHECK (category IN ('airplane','helicopter','simulator')),
        engines TEXT CHECK (engines IN ('single','multi')),
        engine_type TEXT,
        notes TEXT
    )""")
    if not existed:
        conn.executemany("INSERT INTO aircraft_types (code, make, model, category, engines, engine_type) VALUES (?,?,?,?,?,?)",
                         SEED_TYPES)
    conn.commit(); conn.close()


def find_aircraft(registration):
    registration = (registration or "").strip().upper()
    if not registration:
        return None
    conn = aircrafts_connect()
    row = conn.execute("SELECT * FROM aircrafts WHERE registration=?", (registration,)).fetchone()
    conn.close(); return row


def list_aircrafts():
    conn = aircrafts_connect()
    rows = conn.execute("SELECT * FROM aircrafts ORDER BY registration").fetchall()
    conn.close(); return rows


def get_aircraft(aircraft_id):
    conn = aircrafts_connect()
    row = conn.execute("SELECT * FROM aircrafts WHERE id=?", (aircraft_id,)).fetchone()
    conn.close(); return row


def save_aircraft(aircraft_id, form):
    registration = (_clean(form.get("registration")) or "").upper()
    if not registration:
        raise ValueError("Immatriculation obligatoire.")
    engine = _clean(form.get("engine"))
    if engine not in (None,) + ENGINES:
        raise ValueError("Type de moteur invalide.")
    values = (registration, _clean(form.get("type")), engine, _clean(form.get("notes")))
    conn = aircrafts_connect()
    try:
        if aircraft_id:
            conn.execute("UPDATE aircrafts SET registration=?, type=?, engine=?, notes=? WHERE id=?", values + (aircraft_id,))
        else:
            conn.execute("INSERT INTO aircrafts (registration, type, engine, notes) VALUES (?,?,?,?)", values)
        conn.commit()
    except Exception as exc:
        conn.rollback()
        if "UNIQUE" in str(exc):
            raise ValueError("Cette immatriculation existe déjà.")
        raise
    finally:
        conn.close()


def delete_aircraft(aircraft_id):
    conn = aircrafts_connect()
    conn.execute("DELETE FROM aircrafts WHERE id=?", (aircraft_id,))
    conn.commit(); conn.close()


# ---------------- aircraft types (global: one row per type, never per registration) ----------------

def find_aircraft_type(code):
    code = (code or "").strip().upper()
    if not code:
        return None
    conn = aircrafts_connect()
    row = conn.execute("SELECT * FROM aircraft_types WHERE UPPER(code)=?", (code,)).fetchone()
    conn.close(); return row


def list_aircraft_types():
    conn = aircrafts_connect()
    rows = conn.execute("SELECT * FROM aircraft_types ORDER BY category, code").fetchall()
    conn.close(); return rows


def get_aircraft_type(type_id):
    conn = aircrafts_connect()
    row = conn.execute("SELECT * FROM aircraft_types WHERE id=?", (type_id,)).fetchone()
    conn.close(); return row


def save_aircraft_type(type_id, form):
    code = (_clean(form.get("code")) or "").upper()
    if not code:
        raise ValueError("Code du type obligatoire.")
    category, engines, engine_type = _clean(form.get("category")), _clean(form.get("engines")), _clean(form.get("engine_type"))
    if category not in (None,) + tuple(c for c, _ in CATEGORIES):
        raise ValueError("Catégorie invalide.")
    if engines not in (None,) + ENGINES:
        raise ValueError("Nombre de moteurs invalide.")
    if engine_type not in (None,) + tuple(c for c, _ in ENGINE_TYPES):
        raise ValueError("Type de moteur invalide.")
    values = (code, _clean(form.get("make")), _clean(form.get("model")), category, engines, engine_type, _clean(form.get("notes")))
    conn = aircrafts_connect()
    try:
        if type_id:
            conn.execute("UPDATE aircraft_types SET code=?, make=?, model=?, category=?, engines=?, engine_type=?, notes=? WHERE id=?",
                         values + (type_id,))
        else:
            conn.execute("INSERT INTO aircraft_types (code, make, model, category, engines, engine_type, notes) VALUES (?,?,?,?,?,?,?)", values)
        conn.commit()
    except Exception as exc:
        conn.rollback()
        if "UNIQUE" in str(exc):
            raise ValueError("Ce type existe déjà.")
        raise
    finally:
        conn.close()


def delete_aircraft_type(type_id):
    conn = aircrafts_connect()
    conn.execute("DELETE FROM aircraft_types WHERE id=?", (type_id,))
    conn.commit(); conn.close()


# ---------------- Excel export ----------------

def _xlsx(sheet, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for index, header in enumerate(headers, 1):
        longest = max([len(str(header))] + [len(str(r[index - 1] if r[index - 1] is not None else "")) for r in rows])
        ws.column_dimensions[get_column_letter(index)].width = min(max(12, longest + 2), 60)
    ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_aircraft_types_xlsx():
    categories, engine_types = dict(CATEGORIES), dict(ENGINE_TYPES)
    rows = [(t["code"], t["make"], t["model"], categories.get(t["category"], ""), ENGINE_LABELS.get(t["engines"], ""),
             engine_types.get(t["engine_type"], ""), t["notes"]) for t in list_aircraft_types()]
    return _xlsx("Types d'appareils", ["Code", "Marque", "Modèle", "Catégorie", "Moteurs", "Type de moteur", "Notes"], rows)


def export_airports_xlsx():
    labels = {"icao_code": "ICAO", "iata_code": "IATA", "full_name": "Nom", "country": "Pays",
              "latitude": "Latitude", "longitude": "Longitude"}
    conn = airport_connect()
    cursor = conn.execute("SELECT * FROM airports ORDER BY COALESCE(icao_code, iata_code, '')")
    columns = [d[0] for d in cursor.description]
    keep = [c for c in columns if c != "id"]
    rows = [tuple(r[c] for c in keep) for r in cursor.fetchall()]
    conn.close()
    return _xlsx("Aéroports", [labels.get(c, c) for c in keep], rows)


# ---------------- Excel import ----------------

AIRPORT_COLUMNS = {
    "icao_code": ("icao", "icao_code"), "iata_code": ("iata", "iata_code"),
    "full_name": ("nom", "full_name", "name"), "country": ("pays", "country"),
    "latitude": ("latitude", "lat"), "longitude": ("longitude", "lon", "lng"),
    "elevation": ("elevation", "altitude"), "timezone": ("timezone", "fuseau horaire"), "notes": ("notes",)}
AIRCRAFT_COLUMNS = {
    "code": ("code", "type"), "make": ("marque", "make"), "model": ("modele", "model"),
    "category": ("categorie", "category"), "engines": ("moteurs", "engines"),
    "engine_type": ("type de moteur", "engine_type"), "notes": ("notes",)}


def _hdr(value):
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", text.lower())


def _file_rows(filename, data, aliases):
    headers, rows = _read_table(filename, data)
    lookup = {_hdr(a): f for f, names in aliases.items() for a in names}
    cols = {i: lookup[_hdr(h)] for i, h in enumerate(headers) if _hdr(h) in lookup}
    out = []
    for values in rows:
        d = {}
        for i, field in cols.items():
            if i < len(values) and field not in d and _clean(values[i]) is not None:
                d[field] = _clean(values[i])
        if d:
            out.append(d)
    return out, set(cols.values())


def _code_for(value, pairs):
    key = _hdr(value)
    for code, label in pairs:
        if key in (_hdr(code), _hdr(label)):
            return code
    return None


def import_airports(filename, data):
    rows, present = _file_rows(filename, data, AIRPORT_COLUMNS)
    if not ({"icao_code", "iata_code"} & present):
        raise FileImportError("Colonne ICAO ou IATA introuvable dans le fichier.")
    added = updated = skipped = 0
    conn = airport_connect()
    try:
        for d in rows:
            icao = d.get("icao_code", "").upper() or None
            iata = d.get("iata_code", "").upper() or None
            if not (icao or iata):
                skipped += 1; continue
            try:
                lat = _float(d.get("latitude"), "Latitude", -90, 90, required=False)
                lon = _float(d.get("longitude"), "Longitude", -180, 180, required=False)
                elev = _float(d.get("elevation"), "Altitude", -1000, 20000, required=False)
            except ValueError:
                skipped += 1; continue
            new = {"icao_code": icao, "iata_code": iata, "full_name": d.get("full_name"), "country": d.get("country"),
                   "latitude": lat, "longitude": lon, "elevation": elev, "timezone": d.get("timezone"), "notes": d.get("notes")}
            hits = [r for r in (
                conn.execute("SELECT * FROM airports WHERE icao_code=?", (icao,)).fetchone() if icao else None,
                conn.execute("SELECT * FROM airports WHERE iata_code=?", (iata,)).fetchone() if iata else None) if r]
            if len({r["id"] for r in hits}) > 1:
                skipped += 1; continue
            try:
                if hits:
                    gaps = {k: v for k, v in new.items() if v is not None and hits[0][k] in (None, "")}
                    if gaps:
                        conn.execute(f"UPDATE airports SET {', '.join(k + '=?' for k in gaps)} WHERE id=?",
                                     (*gaps.values(), hits[0]["id"]))
                        updated += 1
                elif lat is not None and lon is not None:
                    conn.execute(f"INSERT INTO airports ({', '.join(new)}) VALUES ({','.join('?' * len(new))})", tuple(new.values()))
                    added += 1
                else:
                    skipped += 1
            except sqlite3.IntegrityError:
                skipped += 1
        conn.commit()
    finally:
        conn.close()
    return added, updated, skipped


def import_aircraft_types(filename, data):
    rows, present = _file_rows(filename, data, AIRCRAFT_COLUMNS)
    if "code" not in present:
        raise FileImportError("Colonne « Code » introuvable dans le fichier.")
    added = updated = skipped = 0
    conn = aircrafts_connect()
    try:
        for d in rows:
            code = d.get("code", "").upper()
            if not code:
                skipped += 1; continue
            new = {"code": code, "make": d.get("make"), "model": d.get("model"),
                   "category": _code_for(d.get("category"), CATEGORIES),
                   "engines": _code_for(d.get("engines"), tuple(ENGINE_LABELS.items())),
                   "engine_type": _code_for(d.get("engine_type"), ENGINE_TYPES), "notes": d.get("notes")}
            row = conn.execute("SELECT * FROM aircraft_types WHERE UPPER(code)=?", (code,)).fetchone()
            try:
                if row:
                    gaps = {k: v for k, v in new.items() if v is not None and row[k] in (None, "")}
                    if gaps:
                        conn.execute(f"UPDATE aircraft_types SET {', '.join(k + '=?' for k in gaps)} WHERE id=?",
                                     (*gaps.values(), row["id"]))
                        updated += 1
                else:
                    conn.execute(f"INSERT INTO aircraft_types ({', '.join(new)}) VALUES ({','.join('?' * len(new))})", tuple(new.values()))
                    added += 1
            except sqlite3.IntegrityError:
                skipped += 1
        conn.commit()
    finally:
        conn.close()
    return added, updated, skipped
