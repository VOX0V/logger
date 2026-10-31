"""Régressions de l'audit 0.2.1 : perte de données au renommage / à l'import, migration des comptes, fusion de vols,
export Excel, mise en page du logbook, YAML très imbriqué."""
import io
import sqlite3
from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook, load_workbook

from conftest import REFERER

CSV_HEADS = ["year", "month", "day", "type", "registration", "copilot", "departure", "arrival", "remarks"]


def _columns(app):
    from app.config import configurable_columns, load_config
    with app.app_context():
        return {c["column"]: c for c in configurable_columns(load_config("admin"))}


def _csv(app, rows):
    cols = _columns(app)
    lines = [",".join(cols[h]["import_rules"][0] for h in CSV_HEADS)]
    lines += [",".join(str(x) for x in row) for row in rows]
    return "\n".join(lines).encode()


def _import(web, data, name="vols.csv"):
    r = web.c.post("/import", data={"files": (io.BytesIO(data), name), "csrf_token": web.token()},
                   headers=REFERER, follow_redirects=True)
    assert r.status_code == 200


def _users_db(app):
    return sqlite3.connect(Path(app.config["USERS_DIR"]) / "admin" / "admin_data.db")


FIRST = [[2024, 5, 3, "B738", "C-ABCD", "Jones", "CYUL", "CYYZ", "AC100 test"]]
SECOND = [[2024, 6, 10, "C172", "C-WXYZ", "Brown", "CYHU", "CYQB", "AC200 autre"]]


def test_renaming_display_name_keeps_column_and_data(app, admin):
    _import(admin, _csv(app, FIRST))
    cop = _columns(app)["copilot"]
    r = admin.post(f"/settings/category/{cop['position']}/edit",
                   {"display_name": "co pilot", "data_type": "text", "nullable": "1",
                    "visible": "1", "editable": "1", "import_rules": "\n".join(cop["import_rules"])})
    assert r.status_code == 302
    _import(admin, _csv(app, SECOND))
    db = _users_db(app)
    names = [row[1] for row in db.execute("PRAGMA table_info(users)")]
    assert "copilot" in names and "co_pilot" not in names
    assert ("Jones",) in db.execute("SELECT copilot FROM users").fetchall()
    db.close()
    assert _columns(app)["copilot"]["display_name"] == "co pilot"


def test_import_never_drops_a_column_missing_from_the_config(app, admin):
    from app.config import load_config, save_config
    first, second = _csv(app, FIRST), _csv(app, SECOND)
    _import(admin, first)
    with app.app_context():
        cfg = load_config("admin")
        cfg["database"]["columns"] = [c for c in cfg["database"]["columns"] if c["column"] != "copilot"]
        save_config(cfg, "admin")
    _import(admin, second)
    db = _users_db(app)
    assert "copilot" in [row[1] for row in db.execute("PRAGMA table_info(users)")]
    assert ("Jones",) in db.execute("SELECT copilot FROM users").fetchall()
    db.close()


def test_migration_refuses_accounts_differing_only_by_case(make_app):
    from app import SetupError
    from app.accounts import init_accounts_db
    from app.db import accounts_connect
    app = make_app()
    with app.app_context():
        conn = accounts_connect()
        assert "idx_accounts_username_nocase" in [r[1] for r in conn.execute("PRAGMA index_list(accounts)")]
        conn.execute("DROP INDEX idx_accounts_username_nocase")
        conn.execute("INSERT INTO accounts (username, password_hash, role) VALUES ('Admin', 'x', 'user')")
        conn.commit(); conn.close()
        with pytest.raises(SetupError, match="admin"):
            init_accounts_db()


def test_merge_never_mixes_up_flights_whatever_the_order(app):
    from app.merge import merge_context, plan_merge
    cols = ["year", "month", "day", "type", "registration", "departure", "arrival", "remarks", "flight_number", "block_time"]
    existing = [{"id": 1, "year": 2024, "month": 5, "day": 3, "type": "B738", "registration": None, "departure": None,
                 "arrival": None, "remarks": "AC100 YUL-YYZ", "flight_number": "AC100", "block_time": "1.2"}]
    f100 = {"year": 2024, "month": 5, "day": 3, "type": "B738", "remarks": "AC100 YUL-YYZ", "flight_number": "AC100",
            "block_time": "1.2", "registration": "C-FAAA"}
    f200 = {"year": 2024, "month": 5, "day": 3, "type": "B738", "remarks": "AC200 YYZ-YUL", "flight_number": "AC200",
            "block_time": "1.2", "registration": "C-FBBB"}
    with app.app_context():
        for incoming in ([f100, f200], [f200, f100]):
            inserts, updates, stats = plan_merge([dict(e) for e in existing], [dict(x) for x in incoming], cols, merge_context())
            assert [i["flight_number"] for i in inserts] == ["AC200"]
            assert updates == [(1, {"registration": "C-FAAA"})]
            assert stats["disagreements"] == 0


def _sheet(lines=()):
    wb = Workbook(); ws = wb.active; ws.title = "2024"
    ws.cell(31, 8).value = "GRAND TOTAL"
    for i, (month, day, remarks, reg) in enumerate(lines, 1):
        ws.cell(i, 2).value, ws.cell(i, 3).value = month, day
        ws.cell(i, 10).value, ws.cell(i, 5).value = remarks, reg
    return wb


def test_export_adds_a_flight_when_the_sheet_line_is_another_aircraft():
    from app.export import update_workbook
    rows = [{"date": "2024-05-03", "remarks": "", "aircraft_type": "C172", "registration": "C-ZZZZ", "single_engine_pic_day": "1.0"}]
    report = update_workbook(_sheet([(5, 3, None, "C-AAAA")]), rows)["sheets"][0]
    assert report["added"] == 1 and report["present"] == 0


def test_export_twice_adds_nothing_new():
    from app.export import update_workbook
    rows = [{"date": "2024-05-03", "remarks": "AC100 YUL-YYZ", "aircraft_type": "B738", "registration": "C-FAAA", "multi_engine_pic_day": "1.2"},
            {"date": "2024-05-03", "remarks": "AC100 YYZ-YUL", "aircraft_type": "B738", "registration": "C-FAAA", "multi_engine_pic_day": "1.3"}]
    wb = _sheet()
    assert update_workbook(wb, rows)["sheets"][0]["added"] == 2
    buffer = io.BytesIO(); wb.save(buffer); buffer.seek(0)
    second = update_workbook(load_workbook(buffer), rows)["sheets"][0]
    assert second["added"] == 0 and second["present"] == 2


BLOCK = lambda n, color: {"name": f"b{n}", "rows": 10, "color": color}


@pytest.mark.parametrize("blocks", [[], [BLOCK(1, "#f1dddd")], [BLOCK(1, "#f1dddd"), BLOCK(2, "#dce7f2")]])
def test_logbook_page_survives_fewer_than_three_color_blocks(app, admin, blocks):
    path = Path(app.config["USERS_DIR"]) / "admin" / "admin_layout.yml"
    layout = yaml.safe_load(path.read_text(encoding="utf-8"))
    layout["styles"]["data_blocks"] = blocks
    path.write_text(yaml.safe_dump(layout, allow_unicode=True, sort_keys=False), encoding="utf-8")
    assert admin.get("/logbook").status_code == 200


def test_yaml_validation_reports_absurd_nesting_instead_of_crashing():
    from app.yamledit import validate
    deep = "database:\n  columns: " + "[" * 20000 + "]" * 20000
    assert "nested" in validate("data", deep)
